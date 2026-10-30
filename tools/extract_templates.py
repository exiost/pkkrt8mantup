#!/usr/bin/env python3
"""
Membaca DESAIN lembar laporan dari file Excel asli (lebar kolom, tinggi baris, sel gabungan, garis, huruf,
teks tetap) dan menyimpannya sebagai src/sheet/templates.json. Data warga TIDAK ikut tersimpan:
semua sel berumus dan sel isian dikosongkan; aplikasi mengisinya sendiri.

Pemakaian:
    python3 tools/extract_templates.py /path/Data_Warga.xlsm src/sheet/templates.json
"""
import json, re, sys, colorsys, xml.etree.ElementTree as ET
import openpyxl
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI, range_boundaries
from openpyxl.styles.colors import COLOR_INDEX

# id, nama sheet, area cetak, regions [(id, awal, akhir)], sel isian dinamis (dikosongkan)
CFG = [
  dict(id='penduduk', sheet='DATA PENDUDUK 15-59', area='A1:G174', regions=[('main', 6, 174)]),
  dict(id='kampungkb', sheet='INVENTARISASI KAMPUNG KB', area='A1:P46', regions=[], dyn=['A2', 'A3']),
  dict(id='pencatatan', sheet='PENCATATAN IBU HAMIL,KB,PUS', area='A1:J23', regions=[('hamil', 3, 7), ('kb', 11, 15), ('pus', 19, 23)]),
  dict(id='wus', sheet='DATA WUS PER BULAN', area='A1:F48', regions=[('main', 5, 48)], dyn=['A2']),
  dict(id='pus', sheet='DATA PUS PER BULAN ', area='A1:I44', regions=[('main', 5, 44)], dyn=['A2']),
  dict(id='tppkk', sheet='data warga tp pkk', area='A1:Q61', regions=[], dyn=['E5']),
  dict(id='datakeluarga', sheet='data keluarga', area='A1:R46', regions=[('m', 22, 29)]),
  dict(id='catatankeluarga', sheet='CATATAN KELUARGA', area='A1:S31', regions=[('main', 13, 31)]),
  dict(id='rekapcatatan', sheet='rekap catatan data dan kegiatan', area='A1:AD31', regions=[('main', 15, 30)], dyn=['U6']),
  dict(id='rekapibuhamil', sheet='REKAPITULASI DATA IBU HAMIL', area='A1:P28', regions=[('main', 17, 25)], dyn=['K4', 'K5']),
  dict(id='htirt', sheet='HT_IRT', area='A1:D37', regions=[('pk', 5, 22), ('irt', 27, 36)], dyn=['B2']),
  dict(id='rekapdawis', sheet='REKAP CAT DATA KEG DAWIS 1', area='A1:AD42', regions=[('main', 16, 41)], dyn=['V3', 'V7']),
  dict(id='ibuhamildawis', sheet='rekap ibu hamil dawis 1', area='A1:Q38', regions=[('main', 15, 27)], dyn=['L3', 'A8', 'A9']),
  dict(id='registerlahir', sheet='bayi lahir ', area='A1:U3', regions=[('main', 2, 3)]),
]

wb = openpyxl.load_workbook(sys.argv[1])
out_path = sys.argv[2]

# ---------- warna tema ----------
theme = []
try:
    root = ET.fromstring(wb.loaded_theme)
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    cs = root.find('.//a:clrScheme', ns)
    order = ['lt1', 'dk1', 'lt2', 'dk2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5', 'accent6', 'hlink', 'folHlink']
    for n in order:
        el = cs.find('a:' + n, ns)[0]
        theme.append(el.attrib.get('val') or el.attrib.get('lastClr'))
except Exception:
    theme = ['FFFFFF', '000000', 'EEECE1', '1F497D', '4F81BD', 'C0504D', '9BBB59', '8064A2', '4BACC6', 'F79646', '0000FF', '800080']

def tint(hexrgb, t):
    if not t: return hexrgb
    r, g, b = [int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + t) if t < 0 else l * (1 - t) + t
    r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0), 1), s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))

def color(c, default=None):
    if c is None: return default
    try:
        if c.type == 'rgb' and isinstance(c.rgb, str): return c.rgb[-6:]
        if c.type == 'theme': return tint(theme[c.theme] if c.theme < len(theme) else '000000', c.tint)
        if c.type == 'indexed':
            v = COLOR_INDEX[c.indexed] if c.indexed < len(COLOR_INDEX) else '00000000'
            return v[-6:]
    except Exception:
        pass
    return default

BORDER = {'thin': 1, 'hair': 1, 'dotted': 1, 'dashed': 1, 'dashDot': 1, 'dashDotDot': 1, 'medium': 2, 'mediumDashed': 2,
          'mediumDashDot': 2, 'mediumDashDotDot': 2, 'slantDashDot': 2, 'thick': 3, 'double': 3}
LINE = {'hair': 'dotted', 'dotted': 'dotted', 'dashed': 'dashed', 'mediumDashed': 'dashed', 'dashDot': 'dashed',
        'dashDotDot': 'dashed', 'mediumDashDot': 'dashed', 'mediumDashDotDot': 'dashed', 'slantDashDot': 'dashed', 'double': 'double'}

def bd(side):
    if side is None or not side.style: return None
    return [BORDER.get(side.style, 1), LINE.get(side.style, 'solid'), color(side.color, '000000')]

styles, style_idx = [], {}
def style_of(cell, edges):
    f, a = cell.font, cell.alignment
    fill = None
    if cell.fill and cell.fill.fill_type == 'solid':
        fill = color(cell.fill.fgColor)
        if fill in ('FFFFFF', None) or cell.fill.fgColor.type == 'indexed' and cell.fill.fgColor.indexed in (64,): fill = None
    s = {
        'f': f.name or 'Arial', 'z': float(f.sz or 10),
        'b': 1 if f.b else 0, 'i': 1 if f.i else 0, 'u': 1 if f.u else 0, 'fc': color(f.color, '000000'),
        'h': a.horizontal or 'general', 'v': a.vertical or 'bottom', 'w': 1 if a.wrap_text else 0, 'in': a.indent or 0,
        'rot': a.text_rotation or 0, 'bg': fill, 'fmt': cell.number_format,
        'bt': edges[0], 'br': edges[1], 'bb': edges[2], 'bl': edges[3],
    }
    s = {k: v for k, v in s.items() if v not in (None, 0, False, 'general', 'General') or k in ('f', 'z')}
    if s.get('fc') == '000000': s.pop('fc')
    if s.get('v') == 'bottom': s.pop('v')
    key = json.dumps(s, sort_keys=True)
    if key not in style_idx:
        style_idx[key] = len(styles); styles.append(s)
    return style_idx[key]

def date_fmt(fmt):
    return 'd' if re.search(r'[dmy]', fmt or '', re.I) and not re.search(r'[0#]', fmt or '') and fmt != 'General' else None

result = {}
for cfg in CFG:
    ws = wb[cfg['sheet']]
    c1, r1, c2, r2 = range_boundaries(cfg['area'])
    merges = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        if m.min_row < r1 or m.min_col < c1 or m.max_row > r2 or m.max_col > c2: continue
        merges[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
        for rr in range(m.min_row, m.max_row + 1):
            for cc in range(m.min_col, m.max_col + 1):
                if (rr, cc) != (m.min_row, m.min_col): covered.add((rr, cc))
    # lebar kolom (px)
    widths = {}
    for k, d in ws.column_dimensions.items():
        for ci in range(d.min or CI(k), (d.max or CI(k)) + 1):
            widths[ci] = 0 if d.hidden else d.width
    cols = []
    for ci in range(c1, c2 + 1):
        w = widths.get(ci)
        cols.append(round((w if w is not None else 8.43) * 7) if w != 0 else 0)
    dfh = ws.sheet_format.defaultRowHeight or 12.75
    regions = []
    skip_rows = set()
    for rid, a, b in cfg['regions']:
        regions.append(dict(id=rid, start=a, end=b, size=b - a + 1))
        skip_rows.update(range(a + 1, b + 1))   # hanya baris pertama region (prototipe) yang disimpan
    region_rows = {r for _, a, b in cfg['regions'] for r in range(a, b + 1)}
    dyn = set(cfg.get('dyn', []))
    rows = []
    for r in range(r1, r2 + 1):
        if r in skip_rows: continue
        h = ws.row_dimensions[r].height or dfh
        cells = []
        for c in range(c1, c2 + 1):
            if (r, c) in covered: continue
            cell = ws.cell(r, c)
            rs, cs = merges.get((r, c), (1, 1))
            # Garis dihitung per sisi sel (tabel HTML memakai border terpisah agar tinggi baris persis):
            # sisi kanan/bawah = gabungan garis sel ini & sel tetangga; sisi atas/kiri hanya di tepi area cetak.
            def B(rr, cc, side):
                if not (r1 <= rr <= r2 and c1 <= cc <= c2): return None
                return bd(getattr(ws.cell(rr, cc).border, side))
            def strongest(xs):
                best = None
                for x in xs:
                    if x and (best is None or x[0] > best[0]): best = x
                return best
            top = strongest([B(r, cc, 'top') for cc in range(c, c + cs)]) if r == r1 else None
            lef = strongest([B(rr, c, 'left') for rr in range(r, r + rs)]) if c == c1 else None
            rig = strongest([x for rr in range(r, r + rs) for x in (B(rr, c + cs - 1, 'right'), B(rr, c + cs, 'left'))])
            bot = strongest([x for cc in range(c, c + cs) for x in (B(r + rs - 1, cc, 'bottom'), B(r + rs, cc, 'top'))])
            v = cell.value
            addr = cell.coordinate
            is_f = isinstance(v, str) and v.startswith('=')
            if is_f or addr in dyn or (r in region_rows):
                v = None
            if isinstance(v, str) and not v.strip(): v = v if v else None
            if v is not None and not isinstance(v, (str, int, float)): v = str(v)
            has_style = any([top, bot, lef, rig]) or (cell.fill and cell.fill.fill_type == 'solid' and color(cell.fill.fgColor) not in (None, 'FFFFFF'))
            if v is None and not has_style and (rs, cs) == (1, 1) and not is_f and addr not in dyn and r not in region_rows: continue
            if v is None and not has_style and (rs, cs) == (1, 1) and r in region_rows and not any(x for x in [top, bot, lef, rig]): 
                # sel isian region tanpa garis: tetap dicatat agar rata kiri/kanan benar
                pass
            st = style_of(cell, [top, rig, bot, lef])
            item = {'c': c - c1, 's': st}
            if v is not None: item['v'] = v
            if rs > 1: item['rs'] = rs
            if cs > 1: item['cs'] = cs
            if is_f or addr in dyn or r in region_rows: item['d'] = 1
            cells.append(item)
        rows.append({'r': r, 'h': round(h * 96 / 72, 2), 'c': cells})
    pg, m = ws.page_setup, ws.page_margins
    result[cfg['id']] = {
        'sheet': cfg['sheet'], 'cols': cols, 'rows': rows, 'regions': regions,
        'page': {'orient': pg.orientation or 'portrait', 'paper': pg.paperSize or 9, 'scale': pg.scale or 100,
                 'ml': m.left, 'mr': m.right, 'mt': m.top, 'mb': m.bottom, 'hc': 1 if ws.print_options.horizontalCentered else 0},
    }
out = {'styles': styles, 'sheets': result}
json.dump(out, open(out_path, 'w'), ensure_ascii=False, separators=(',', ':'))
print('ok', out_path, len(styles), 'styles;', {k: len(v['rows']) for k, v in result.items()})
